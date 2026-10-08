from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "AssetTrustMonitor_Development_Log_1.4_to_1.4.4.docx"


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_entry(doc, title, found, changes, result):
    doc.add_heading(title, level=2)
    paragraph = doc.add_paragraph()
    paragraph.add_run("What I found: ").bold = True
    paragraph.add_run(found)
    paragraph = doc.add_paragraph()
    paragraph.add_run("What I changed: ").bold = True
    paragraph.add_run(changes)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Result: ").bold = True
    paragraph.add_run(result)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
for style_name, size, color in (
    ("Title", 25, RGBColor(35, 35, 35)),
    ("Heading 1", 17, RGBColor(24, 105, 82)),
    ("Heading 2", 12.5, RGBColor(35, 35, 35)),
):
    styles[style_name].font.name = "Arial"
    styles[style_name].font.size = Pt(size)
    styles[style_name].font.color.rgb = color

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("AssetTrustMonitor Development Log")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Versions 1.4 through 1.4.4\n").bold = True
subtitle.add_run("Notes from testing, bug fixes, and release preparation")

doc.add_paragraph(
    "This is a working development log rather than a formal report. It records the problems I ran into, "
    "how I tracked them down, and the main code or packaging changes made while moving from version 1.4 to 1.4.4."
)

doc.add_heading("Version 1.4 - starting point", level=1)
doc.add_paragraph(
    "Version 1.4 already monitored asset folders, processed Roblox model files, worked with Rojo, and displayed results in a desktop interface. "
    "Most later work came from testing the same workflow on different Windows computers, testing on macOS, and repeatedly importing real .rbxm files."
)

doc.add_heading("Version 1.4.1 - connection and monitoring fixes", level=1)
add_entry(
    doc,
    "Rojo stopped or could not connect",
    "The application could start before the installed Rojo tool was ready, or an older rojo serve process could still own the port.",
    "The startup path was reorganized to check the tool, stop only the process owned by AssetTrustMonitor, start rojo serve with an explicit project file and working directory, and verify the port before continuing.",
    "Rojo startup became more predictable, and failures now show a useful message instead of silently stopping."
)
add_entry(
    doc,
    "Assets reached the folder but not Roblox Studio",
    "Copying a file into the monitored Assets folder only updates the filesystem. Studio still needs an active Rojo connection and a matching project mapping before the object appears under Workspace.",
    "The monitor and bridge status were separated so the interface could show whether the file was imported, whether Rojo was running, and whether Studio had actually synchronized. Startup ordering was also tightened.",
    "It became easier to distinguish a successful file import from a successful Studio sync."
)
add_entry(
    doc,
    "Logging slowdown and slow shutdown",
    "Frequent UI log updates and an unbounded queue could make long sessions sluggish. File-monitor threads could also wait too long during shutdown.",
    "Logging was routed through a bounded queue, repeated updates were reduced, UTF-8 rotating log files were used, and monitor waits were made interruptible so stop events are checked sooner.",
    "The UI stays responsive during busy imports and closes faster."
)
add_entry(
    doc,
    "Different virtual environments on different computers",
    "Hard-coded Python or environment paths did not transfer cleanly to another computer.",
    "Launch and bootstrap code now resolves paths relative to the application folder, creates or validates a private environment locally, and reports setup progress in the splash window.",
    "The package is less dependent on the original development machine."
)

doc.add_heading("Version 1.4.2 - safety and code cleanup", level=1)
add_entry(
    doc,
    "Windows path-too-long errors",
    "The original nested release folders made files such as AssetTrustMirrorPlugin.lua exceed Windows extraction limits.",
    "Release staging was moved to a shorter folder name, duplicate nesting and development-only files were removed, and ZIP contents were checked before release.",
    "The package extracts on a second computer without requiring unusually long paths."
)
add_entry(
    doc,
    "Unnecessary shell=True usage",
    "Some subprocess calls used the command shell even though they were starting known executables with fixed arguments.",
    "Commands were changed to argument lists and sent through the shared command runner. Windows hidden-window flags, working directories, timeouts, and captured output were applied where needed.",
    "Commands are easier to validate and no longer depend on an unnecessary shell window."
)
add_entry(
    doc,
    "Broad exception handling",
    "Several try/except blocks hid the reason a tool, file, or UI step failed.",
    "Operational paths gained narrower exception handling and contextual logging. Broad catches were kept only at cleanup, UI, and compatibility boundaries where a third-party failure should not crash the entire program.",
    "Errors are more useful while the program still protects the main interface from unexpected failures."
)
add_entry(
    doc,
    "Possible SQL injection warning",
    "Bandit reported a possible SQL-related finding, but tracing the flagged value showed no database connection and no dynamically constructed SQL query in the runtime path.",
    "The code path was reviewed from input to use. The project uses files, TOML configuration, HTTP checks, and subprocesses in that area, so the warning was documented as a false positive rather than suppressed without review.",
    "No exploitable SQL injection path was found."
)
add_entry(
    doc,
    "Unsafe or malformed .rbxm files stopped the plugin",
    "Some models contained unsupported structures or scripts that caused a large import to fail as one operation.",
    "Validation and error boundaries were added around model processing. Files are handled individually, invalid items are logged or quarantined, and one failed asset does not stop the watcher or the remaining import queue.",
    "Problem assets can be investigated separately without taking down the whole monitoring session."
)

doc.add_heading("Version 1.4.3 - startup and packaging", level=1)
add_entry(
    doc,
    "Startup looked unfinished",
    "Setup messages opened in console windows, and the main interface could appear before initialization was complete.",
    "A small PySide6 splash component was added with a centered banner, simple fade in/out, status text, and staged startup messages. It reuses the existing QApplication and opens the main window once initialization finishes.",
    "First-time setup is visible without exposing command windows."
)
add_entry(
    doc,
    "Console windows appeared behind the banner",
    "Python, Rokit, and Rojo subprocesses inherited a visible console on Windows.",
    "The launcher and subprocess helper were updated to use a windowless executable or Windows no-window creation flags. Command output is captured and forwarded to the splash status instead of printed in a separate terminal.",
    "Normal startup no longer leaves a console behind the application."
)
add_entry(
    doc,
    "Character encoding crash after Studio opened",
    "A Unicode symbol could not be encoded by the active Windows charmap, causing a fatal error while writing output.",
    "User-facing status text was kept Unicode-safe, log files were opened as UTF-8, and captured subprocess output was decoded with an explicit error policy.",
    "The application no longer fails because the Windows console code page cannot represent a symbol."
)
add_entry(
    doc,
    "Launcher and application identity",
    "VBS launchers looked suspicious to users, batch files could flash a terminal, and generic executable packaging sometimes triggered antivirus heuristics.",
    "The release uses a normal application launcher, an AssetTrustMonitor .ico file, consistent executable metadata, and only the files required by the selected operating system. Unused launchers and duplicate processor variants were removed from customer folders.",
    "The package looks and behaves more like a regular desktop application."
)

doc.add_heading("Version 1.4.4 - updater and cross-platform work", level=1)
add_entry(
    doc,
    "Rojo 7.7.1 and Rokit tool updates",
    "A clean computer could install Rokit successfully but attempt rojo serve before the selected Rojo version was installed or active. Tool updates were also hidden in the Debug area.",
    "The release declaration pins Rojo 7.7.1. Version selection and update checks were moved into the normal interface. Setup now installs or activates the requested version, verifies rojo --version, and only then starts the server.",
    "Users can update Rojo without opening Debug, and a success message is not shown until the usable tool is verified."
)
add_entry(
    doc,
    "Retry reported failure after a successful install",
    "One retry branch returned a false value even after installation completed, so the next startup step treated success as failure.",
    "The return value and ordered startup state were corrected. Setup, server launch, port verification, and Studio launch now advance only after the prior stage succeeds.",
    "The misleading 'failed to start Rojo server' message is avoided after a successful repair."
)
add_entry(
    doc,
    "Workspace appeared only after Studio restart",
    "Studio accepted the Rojo sync notification before the project tree was fully refreshed. The filesystem was correct, but the connected Studio session had stale mapping state.",
    "The startup flow now waits for a ready server before Studio sync and gives clearer reconnect guidance. Restarting worked because it forced Studio and the plugin to request the project tree again.",
    "New sessions receive a more consistent Workspace mapping, while reconnect remains the recovery step for a stale Studio session."
)
add_entry(
    doc,
    "macOS startup and Tcl/Tk",
    "Python installations on macOS can use different Tcl/Tk locations, and Windows-only launch logic cannot be reused directly.",
    "Platform checks were added around launcher, process, and Tcl/Tk handling. Windows and macOS releases now expose one launcher each, while processor-specific or platform-specific helpers stay internal only when required.",
    "The public package is easier to understand and avoids presenting incompatible launchers."
)

doc.add_heading("What the review tools found", level=1)
doc.add_heading("Ruff", level=2)
add_bullets(doc, [
    "Highlighted unused or inconsistent imports and general maintainability issues.",
    "Helped identify code paths that could be simplified without removing optional platform-specific imports.",
    "Warnings were reviewed individually instead of applying automatic changes to working compatibility code."
])
doc.add_heading("Bandit", level=2)
add_bullets(doc, [
    "Flagged subprocess calls using shell=True, which led to replacing unnecessary shell commands with argument lists.",
    "Prompted a review of process termination and file operations so they remain limited to application-owned processes and configured folders.",
    "Reported a possible SQL injection issue; tracing the code showed it was a false positive because that path does not execute SQL."
])
doc.add_heading("Memory profiler and long-duration test", level=2)
doc.add_paragraph(
    "The test generated 10,000 synthetic log messages over about 35 seconds and collected 18 samples. The queue stopped at its 2,000-entry limit, "
    "Python traced memory settled near 0.35 MiB with a peak around 0.376 MiB, worker threads returned to the starting count, and process working set moved from about 37.426 MiB to 38.293 MiB. "
    "This short run did not show uncontrolled queue growth, but it is not proof that the program is leak-free. A one-hour and overnight soak test is still recommended for release validation."
)

doc.add_heading("Main files and code areas adjusted", level=1)
add_bullets(doc, [
    "Application entry point: startup order, QApplication ownership, splash-to-main-window transition, and fatal error reporting.",
    "ui/splash/splash_window.py: banner layout, status stages, fade timing, and clean finish behavior.",
    "Subprocess/tool helpers: argument-list commands, hidden-window flags, captured UTF-8 output, timeouts, working directories, and return-value checks.",
    "Rojo service code: selected-version validation, rojo serve project path, port readiness, owned-process shutdown, retry handling, and reconnect status.",
    "File monitor/import code: bounded logging, faster stop checks, per-file exception boundaries, invalid model handling, and continued processing after failures.",
    "Packaging and launcher files: shorter release paths, application icon, one public launcher per OS, and removal of development-only material.",
    "rokit.toml or release tool declaration: Rojo version pinned to 7.7.1 for the 1.4.4 package."
])

doc.add_heading("Remaining checks before calling 1.4.4 final", level=1)
add_bullets(doc, [
    "Test a clean Windows computer with no existing Rokit, Rojo, Python environment, or Studio connection.",
    "Test first launch and second launch separately so installation and normal startup are both covered.",
    "Import small, large, scripted, malformed, and duplicate .rbxm files while watching both the monitor and Studio Workspace.",
    "Disconnect and reconnect the Studio plugin to confirm stale mapping recovery.",
    "Run at least a one-hour memory and CPU soak test; run an overnight test before wider release if possible.",
    "Repeat the startup, Tcl/Tk, launcher, Rojo, and Studio checks on the supported macOS build."
])

doc.add_heading("End note", level=1)
doc.add_paragraph(
    "The biggest lesson from these versions was that a file reaching the Assets folder, Rojo running, and Studio showing the model are three separate states. "
    "Most reliability work in 1.4.4 focuses on checking those states in order and reporting the exact stage that failed."
)

doc.core_properties.title = "AssetTrustMonitor Development Log - Versions 1.4 to 1.4.4"
doc.core_properties.subject = "Development notes, fixes, testing, and tool findings"
doc.core_properties.author = "AssetTrustMonitor Development"
doc.save(OUTPUT)
print(OUTPUT)
